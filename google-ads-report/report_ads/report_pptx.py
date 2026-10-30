"""Report PowerPoint (16:9, Georgia e Calibri,
copertina e chiusura scure, pagine su fondo chiaro, grafici nativi modificabili)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .report_locali import _it, eur, num

TEMI = {
    "rosso": dict(scuro="4A0808", primario="8B0000", sfondo="FBF5F0", testo="2B2422", tenue="7A6F6A", bordo="E2D6CC",
                  oro="C9A227", cover_sub="F0D9D7", cover_txt="E8C8C5", cover_piede="C9A09C", verde="2E7D4F",
                  verde_bg="E7F2EA", crema="FBF1DF", ocra="A86A12", prec="E6D9CF"),
    "verde": dict(scuro="1B3830", primario="2F5D50", sfondo="F6F4EE", testo="23302C", tenue="6E7A75", bordo="D5DCD8",
                  oro="B7791F", cover_sub="D9E6E0", cover_txt="C9D8D1", cover_piede="A9BDB4", verde="2E7D4F",
                  verde_bg="E5F1E8", crema="EFE7D6", ocra="B7791F", prec="DDE3DF"),
}
GIORNI = ["Lunedì", "Martedì", "Mercoledì", "Giovedì", "Venerdì", "Sabato", "Domenica"]


def tema_per(colore: str) -> dict:
    try:
        r, g = int(colore[1:3], 16), int(colore[3:5], 16)
    except (ValueError, TypeError):
        return TEMI["rosso"]
    return TEMI["rosso"] if r > g else TEMI["verde"]


def _pct(x, dec=1):
    return "–" if x is None else f"{_it(x * 100, dec)}%"


def _var(cur, prev):
    if cur is None or not prev:
        return "–"
    d = cur / prev - 1
    if abs(d) < 0.005:
        return "stabile"
    return f"{'+' if d > 0 else '−'}{_it(abs(d) * 100, 0)}%"


class Deck:
    def __init__(self, tema: dict, piede: str):
        self.t = tema
        self.piede = piede
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.layout = self.prs.slide_layouts[6]
        self.n = 0

    # --- primitive
    def slide(self, scura=False):
        s = self.prs.slides.add_slide(self.layout)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string(self.t["scuro"] if scura else self.t["sfondo"])
        self.n += 1
        if not scura:
            self.testo(s, 0.6, 7.0, 9, 0.3, self.piede, size=9, colore=self.t["tenue"])
            self.testo(s, 12.3, 7.0, 0.6, 0.3, str(self.n), size=9, colore=self.t["tenue"], align=PP_ALIGN.RIGHT)
        return s

    def forma(self, s, x, y, w, h, fill, forma=MSO_SHAPE.ROUNDED_RECTANGLE, bordo=None):
        sh = s.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fill)
        if bordo:
            sh.line.color.rgb = RGBColor.from_string(bordo)
            sh.line.width = Pt(0.75)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        if forma == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = 0.06
        return sh

    def testo(self, s, x, y, w, h, testo, size=12, colore=None, bold=False, font="Calibri", align=PP_ALIGN.LEFT,
              italic=False, anchor=MSO_ANCHOR.TOP):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        righe = testo if isinstance(testo, list) else [testo]
        for i, riga in enumerate(righe):
            par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            par.alignment = align
            spec = riga if isinstance(riga, dict) else {"s": riga}
            self._runs(par, spec["s"], spec.get("size", size), spec.get("colore", colore or self.t["testo"]),
                       spec.get("bold", bold), spec.get("font", font), spec.get("italic", italic))
            par.space_after = Pt(spec.get("dopo", 4))
        return tb

    @staticmethod
    def _runs(par, s, size, colore, bold, font, italic):
        for k, pezzo in enumerate(str(s).split("**")):
            if not pezzo:
                continue
            r = par.add_run()
            r.text = pezzo
            r.font.size, r.font.name, r.font.italic = Pt(size), font, italic
            r.font.bold = bold or (k % 2 == 1)
            r.font.color.rgb = RGBColor.from_string(colore)

    def titolo(self, s, titolo, sottotitolo=""):
        self.testo(s, 0.6, 0.38, 12.1, 0.8, titolo, size=30, colore=self.t["primario"], bold=True, font="Georgia")
        if sottotitolo:
            self.testo(s, 0.6, 1.12, 12.1, 0.45, sottotitolo, size=14, colore=self.t["tenue"], italic=True)

    def kpi(self, s, x, y, w, h, valore, etichetta, nota="", evidenza=False):
        self.forma(s, x, y, w, h, self.t["verde_bg"] if evidenza else "FFFFFF", bordo=None if evidenza else self.t["bordo"])
        col = self.t["verde"] if evidenza else self.t["primario"]
        self.testo(s, x, y + 0.22, w, 0.75, valore, size=34, colore=col, bold=True, font="Georgia", align=PP_ALIGN.CENTER)
        self.testo(s, x + 0.15, y + 0.98, w - 0.3, 0.35, etichetta, size=13, bold=True, align=PP_ALIGN.CENTER)
        if nota:
            self.testo(s, x + 0.15, y + 1.32, w - 0.3, 0.5, nota, size=11, colore=self.t["tenue"], align=PP_ALIGN.CENTER)

    def tabella(self, s, x, y, w, righe, larghezze, evidenzia=None, altezza=0.36, size=12):
        n_r, n_c = len(righe), len(righe[0])
        tb = s.shapes.add_table(n_r, n_c, Inches(x), Inches(y), Inches(w), Inches(altezza * n_r)).table
        tot = sum(larghezze)
        for j, lw in enumerate(larghezze):
            tb.columns[j].width = Inches(w * lw / tot)
        tbl_pr = tb._tbl.tblPr
        for attr in ("firstRow", "bandRow"):
            tbl_pr.set(attr, "0")
        for i, riga in enumerate(righe):
            tb.rows[i].height = Inches(altezza)
            for j, val in enumerate(riga):
                c = tb.cell(i, j)
                fill = self.t["primario"] if i == 0 else ((evidenzia or {}).get(i) or "FFFFFF")
                c.fill.solid()
                c.fill.fore_color.rgb = RGBColor.from_string(fill)
                c.margin_left = c.margin_right = Inches(0.1)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = c.text_frame
                par = tf.paragraphs[0]
                par.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
                testo = str(val)
                grassetto = i == 0 or j == 0 and (evidenzia or {}).get(i) is not None
                self._runs(par, testo, size, "FFFFFF" if i == 0 else self.t["testo"], grassetto, "Calibri", False)
                self._bordi(c)
        return tb

    def _bordi(self, cella):
        tc_pr = cella._tc.get_or_add_tcPr()
        for lato in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
            ln = tc_pr.find(qn(lato))
            if ln is None:
                ln = tc_pr.makeelement(qn(lato), {})
                tc_pr.append(ln)
            ln.set("w", "6350")
            for figlio in list(ln):
                ln.remove(figlio)
            fill = ln.makeelement(qn("a:solidFill"), {})
            clr = fill.makeelement(qn("a:srgbClr"), {"val": self.t["bordo"]})
            fill.append(clr)
            ln.append(fill)

    def barre(self, s, x, y, w, h, categorie, serie, orizzontale=False, colori=None, picco=None, legenda=True,
              etichette=True):
        dati = CategoryChartData()
        dati.categories = categorie
        for nome, valori in serie:
            dati.add_series(nome, valori)
        tipo = XL_CHART_TYPE.BAR_CLUSTERED if orizzontale else XL_CHART_TYPE.COLUMN_CLUSTERED
        ch = s.shapes.add_chart(tipo, Inches(x), Inches(y), Inches(w), Inches(h), dati).chart
        ch.has_title = False
        ch.has_legend = legenda and len(serie) > 1
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.TOP
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(10)
            ch.legend.font.color.rgb = RGBColor.from_string(self.t["tenue"])
        va = ch.value_axis
        va.visible = False
        va.has_major_gridlines = False
        ca = ch.category_axis
        ca.tick_labels.font.size = Pt(10)
        ca.tick_labels.font.color.rgb = RGBColor.from_string(self.t["tenue"])
        ca.format.line.fill.background()
        if orizzontale:
            ca.reverse_order = True
        plot = ch.plots[0]
        plot.gap_width = 60
        plot.overlap = -10 if len(serie) > 1 else 0
        colori = colori or [self.t["primario"], self.t["prec"]]
        for i, ser in enumerate(plot.series):
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = RGBColor.from_string(colori[i % len(colori)])
            if picco is not None and i == 0:
                for k, v in enumerate(serie[0][1]):
                    if k not in picco:
                        pt = ser.points[k]
                        pt.format.fill.solid()
                        pt.format.fill.fore_color.rgb = RGBColor.from_string(self.t["prec"])
            if etichette:
                ser.data_labels.show_value = True
                ser.data_labels.number_format = '#,##0'
                ser.data_labels.number_format_is_linked = False
                ser.data_labels.font.size = Pt(9)
                ser.data_labels.font.color.rgb = RGBColor.from_string(self.t["testo"])
                ser.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        return ch

    def salva(self, percorso: Path) -> Path:
        self.prs.save(percorso)
        return percorso


def testi_base(d: dict, x: dict, manuali: dict) -> dict:
    cor, pre = x["corrente"], x["precedente"]
    mese, mese_prec = d["mese"], d["mese_prec"]
    t = {
        "sintesi_sottotitolo": f"Il mese di {mese} a confronto con {mese_prec}",
        "sintesi_titolo": f"{num(cor['conversioni'])} conversioni",
        "sintesi_testo": f"Spesa di {eur(cor['spesa'], 0)} (era {eur(pre['spesa'], 0)}), con {num(cor['indicazioni'])} "
                         f"indicazioni stradali e {num(cor['visite'])} visite in negozio.",
        "confronto_lettura": f"Rispetto a {mese_prec} le indicazioni stradali sono {_var(cor['indicazioni'], pre['indicazioni'])} "
                             f"e il costo per conversione è {eur(cor['costo_conversione'])}.",
        "temi_testo": "",
        "lingue_testo": "",
        "canali_titolo": "Maps e Ricerca",
        "canali_testo": "",
        "prossimi": [[p.split(":", 1)[0], p.split(":", 1)[1].strip() if ":" in p else ""] for p in d["prossimi_passi"]][:4]
        or [[f"Rileggere le visite di {mese} a dato consolidato", f"Il confronto vero si fa a fine {d['mese_succ']}."]],
    }
    if x["lingue"]:
        top = max(x["lingue"], key=lambda g: g["conversioni"])
        t["lingue_testo"] = f"{top['nome']} porta il volume maggiore con {num(top['conversioni'])} conversioni."
    maps = next((c for c in x["canali"] if c["nome"] == "Google Maps"), None)
    if maps:
        tot_ind = sum(c["indicazioni"] for c in x["canali"]) or 1
        t["canali_testo"] = f"Maps porta il {_pct(maps['indicazioni'] / tot_ind, 0)} delle indicazioni stradali."
    return t


def deck_grano(d: dict, x: dict, t: dict, manuali: dict, colore: str, percorso: Path) -> Path:
    tema = tema_per(colore)
    mese, mese_prec, anno = d["mese"], d["mese_prec"], d["anno"]
    cor, pre = x["corrente"], x["precedente"]
    k = Deck(tema, f"{x['nome']}  ·  Report Google Ads  ·  {mese.capitalize()} {anno}")
    T = tema
    stimate = x["stimate"]

    # --- 1 copertina
    s = k.slide(scura=True)
    k.forma(s, 9.4, -2.2, 6.2, 6.2, T["primario"], MSO_SHAPE.OVAL)
    k.forma(s, -1.6, 5.2, 4.0, 4.0, T["primario"], MSO_SHAPE.OVAL)
    k.testo(s, 0.9, 1.6, 8, 0.4, "REPORT GOOGLE ADS" + (f"  ·  N.{d['numero']}" if d["numero"] else ""), size=14,
            colore=T["oro"], bold=True)
    k.testo(s, 0.9, 2.1, 11.5, 1.2, x["nome"], size=50 if len(x["nome"]) > 18 else 58, colore="FFFFFF", bold=True,
            font="Georgia")
    k.testo(s, 0.9, 3.35, 8, 0.7, f"{mese.capitalize()} {anno}", size=30, colore=T["cover_sub"], font="Georgia")
    linea = s.shapes.add_connector(1, Inches(0.92), Inches(4.35), Inches(4.12), Inches(4.35))
    linea.line.color.rgb = RGBColor.from_string(T["oro"])
    linea.line.width = Pt(2)
    riga = "  ·  ".join(v for v in (x["descrizione"], f"Confronto con {mese_prec}",
                                     f"Focus: {d['focus']}" if d["focus"] else "") if v)
    k.testo(s, 0.9, 4.55, 11, 0.4, riga, size=15, colore=T["cover_txt"])
    k.testo(s, 0.9, 6.6, 8, 0.35, "Francesco Gizzi · Growth & Marketing", size=12, colore=T["cover_piede"])

    # --- 2 sintesi
    s = k.slide()
    k.titolo(s, f"{mese.capitalize()} in sintesi", t["sintesi_sottotitolo"])
    carte = [
        (num(cor["indicazioni"]), "Indicazioni stradali", f"erano {num(pre['indicazioni'])} a {mese_prec}", False),
        (num(cor["visite"]), "Visite stimate" if stimate else "Visite in negozio",
         f"{_pct(x['quota_stima'], 0)} delle indicazioni" if stimate else "dato provvisorio, vedi slide 3", False),
        (eur(cor["costo_conversione"]), "Costo per conversione", f"era {eur(pre['costo_conversione'])} a {mese_prec}", True),
        (_pct(cor["ctr"], 2), "CTR", f"era {_pct(pre['ctr'], 2)} a {mese_prec}", True),
    ]
    for i, (v, e, nota, ev) in enumerate(carte):
        k.kpi(s, 0.6 + i * 3.1, 1.85, 2.9, 1.95, v, e, nota, ev)
    k.forma(s, 0.6, 4.1, 12.2, 1.9, "FFFFFF", bordo=T["bordo"])
    righe = [{"s": t["sintesi_titolo"], "size": 22, "bold": True, "font": "Georgia", "colore": T["primario"], "dopo": 8},
             {"s": t["sintesi_testo"], "size": 13}]
    if x["giorni_attivi"] < x["giorni_mese"]:
        righe.append({"s": f"La campagna ha lavorato {x['giorni_attivi']} giorni su {x['giorni_mese']}: i confronti per "
                           "giorno sono calcolati sui giorni effettivi.", "size": 11, "italic": True, "colore": T["tenue"]})
    k.testo(s, 0.95, 4.3, 11.5, 1.6, righe)

    # --- 3 visite
    s = k.slide()
    if stimate:
        k.titolo(s, "Come stimiamo le visite", f"Per {mese} il riferimento sono le indicazioni stradali")
        k.kpi(s, 0.6, 1.85, 3.4, 1.9, num(cor["indicazioni"]), "Indicazioni stradali", "dato definitivo")
        k.testo(s, 4.1, 2.45, 1.2, 0.6, f"× {_it(x['quota_stima'] * 100, 0)}%", size=22, colore=T["oro"], bold=True,
                font="Georgia", align=PP_ALIGN.CENTER)
        k.kpi(s, 5.4, 1.85, 3.4, 1.9, num(cor["visite"]), "Visite stimate", "", True)
        k.testo(s, 0.6, 4.2, 12.1, 2.4, [
            {"s": f"**Perché le indicazioni.** Quando la metrica delle visite in negozio non è completa per tutto il mese, "
                  f"come da accordo si usa il {_it(x['quota_stima'] * 100, 0)}% delle richieste di indicazioni stradali.",
             "size": 13, "dopo": 10},
            {"s": f"**Perché il {_it(x['quota_stima'] * 100, 0)}%.** È un valore prudente, misurato sugli altri locali che "
                  "gestiamo dove la metrica è attiva da mesi. Le indicazioni sono un dato definitivo, quindi non serve un conguaglio.",
             "size": 13}])
    else:
        k.titolo(s, f"Le visite di {mese} sono ancora da completare",
                 "Google registra le visite in negozio fino a circa 30 giorni dopo il clic sull'annuncio")
        k.forma(s, 0.6, 1.85, 5.6, 3.3, "FFFFFF", bordo=T["bordo"])
        k.testo(s, 0.9, 2.1, 5, 0.4, "Visite registrate oggi", size=13, colore=T["tenue"])
        k.testo(s, 0.9, 2.6, 5, 1.3, num(cor["visite"]), size=66, colore=T["primario"], bold=True, font="Georgia")
        k.testo(s, 0.9, 4.2, 5, 0.4, f"conteggio parziale al {d['fine'].day + 0} {mese}", size=11, colore=T["tenue"],
                italic=True)
        k.testo(s, 6.25, 2.9, 0.6, 0.8, "→", size=40, colore=T["oro"], align=PP_ALIGN.CENTER)
        k.forma(s, 6.9, 1.85, 5.9, 3.3, T["verde_bg"])
        k.testo(s, 7.2, 2.1, 5.3, 0.4, "Stima a dato consolidato", size=13, colore=T["verde"])
        k.testo(s, 7.2, 2.6, 5.3, 1.3, num(round(cor["visite"] * 1.3, -1)) + "+", size=66, colore=T["verde"], bold=True,
                font="Georgia")
        k.testo(s, 7.2, 4.2, 5.3, 0.4, "almeno il 30% in più, come visto nei mesi scorsi", size=11, colore=T["verde"],
                italic=True)
        k.testo(s, 0.6, 5.45, 12.2, 0.9, f"Ogni mese una parte rilevante delle visite viene attribuita nelle settimane "
                f"successive. Il dato di {mese} va quindi letto come un minimo: il confronto corretto con {mese_prec} si "
                f"potrà fare a fine {d['mese_succ']}.", size=12)

    # --- 4 confronto
    s = k.slide()
    ga, gp = x["giorni_attivi"], x["giorni_prec"]
    k.titolo(s, f"Il confronto con {mese_prec}",
             f"Totali del mese e valore per giorno di campagna attiva ({gp} giorni a {mese_prec}, {ga} a {mese})")

    def per_giorno(c, p):
        if c is None or not p:
            return "–"
        dlt = (c / ga) / (p / gp) - 1
        return "stabile" if abs(dlt) < 0.01 else f"{'+' if dlt > 0 else '−'}{_it(abs(dlt) * 100, 0)}%"

    righe = [["Metrica", mese_prec.capitalize(), mese.capitalize(), "Per giorno attivo"],
             ["Conversioni", num(pre["conversioni"]), num(cor["conversioni"]), per_giorno(cor["conversioni"], pre["conversioni"])],
             ["Indicazioni stradali", num(pre["indicazioni"]), num(cor["indicazioni"]), per_giorno(cor["indicazioni"], pre["indicazioni"])],
             ["Costo per conversione", eur(pre["costo_conversione"]), eur(cor["costo_conversione"]),
              _var(cor["costo_conversione"], pre["costo_conversione"])],
             ["CTR", _pct(pre["ctr"], 2), _pct(cor["ctr"], 2),
              f"{'+' if (cor['ctr'] or 0) >= (pre['ctr'] or 0) else '−'}{_it(abs((cor['ctr'] or 0) - (pre['ctr'] or 0)) * 100, 1)} punti"],
             ["Spesa Google Ads", eur(pre["spesa"], 0), eur(cor["spesa"], 0), _var(cor["spesa"], pre["spesa"]) + " sul mese"],
             ["Impressioni", num(pre["impressioni"]), num(cor["impressioni"]), per_giorno(cor["impressioni"], pre["impressioni"])],
             ["Visite stimate" if stimate else "Visite in negozio", num(pre["visite"]), num(cor["visite"]),
              "stima" if stimate else "dato provvisorio"]]
    k.tabella(s, 0.6, 1.8, 8.2, righe, [3, 1.6, 1.6, 2], evidenzia={1: T["verde_bg"], 2: T["verde_bg"], 7: T["crema"]})
    k.forma(s, 9.2, 1.8, 3.6, 3.4, "FFFFFF", bordo=T["bordo"])
    k.testo(s, 9.45, 2.0, 3.15, 3.1, [{"s": "La lettura", "size": 18, "bold": True, "font": "Georgia",
                                       "colore": T["primario"], "dopo": 8}, {"s": t["confronto_lettura"], "size": 12}])

    # --- 5 temi (dalla mail)
    temi = (manuali.get(x["nome"]) or {}).get("temi")
    if temi:
        s = k.slide()
        k.titolo(s, "L'effetto del lavoro sui temi di ricerca",
                 "Temi che Google considera incrementali, cioè capaci di portare traffico in più, prima e dopo l'aggiornamento")
        k.barre(s, 0.6, 1.7, 8.0, 4.9, temi["lingue"], [("Fine mese", temi["dopo"]), ("Inizio mese", temi["prima"])],
                colori=[T["primario"], T["prec"]])
        k.forma(s, 8.9, 1.8, 3.9, 1.7, T["verde_bg"])
        k.testo(s, 8.9, 1.95, 3.9, 0.8, f"{sum(temi['prima'])} → {sum(temi['dopo'])}", size=36, colore=T["verde"],
                bold=True, font="Georgia", align=PP_ALIGN.CENTER)
        k.testo(s, 9.1, 2.8, 3.5, 0.6, f"temi incrementali su {len(temi['lingue'])} lingue", size=12,
                colore=T["verde"], align=PP_ALIGN.CENTER)
        k.testo(s, 8.9, 3.8, 3.9, 2.8, t["temi_testo"], size=12)

    # --- 6 lingue
    if x["lingue"]:
        s = k.slide()
        k.titolo(s, f"Le lingue a {mese}", f"Conversioni per gruppo di lingua, confronto con {mese_prec}")
        lingue = sorted(x["lingue"], key=lambda g: -g["conversioni"])
        serie = [(mese.capitalize(), [round(g["conversioni"]) for g in lingue])]
        if any(g["conversioni_prec"] for g in lingue):
            serie.append((mese_prec.capitalize(), [round(g["conversioni_prec"] or 0) for g in lingue]))
        k.barre(s, 0.6, 1.7, 7.4, 4.95, [g["nome"] for g in lingue], serie, orizzontale=True)
        righe = [["Lingua", "Variaz.", "Tasso conv."]]
        for g in sorted(lingue, key=lambda g: -((g["conversioni"] / g["conversioni_prec"] - 1) if g["conversioni_prec"] else -9)):
            righe.append([g["nome"], _var(g["conversioni"], g["conversioni_prec"]), _pct(g["tasso"])])
        k.tabella(s, 8.3, 1.75, 4.5, righe, [2, 1.2, 1.4], altezza=0.32, size=11)
        k.testo(s, 8.3, 1.85 + 0.32 * len(righe) + 0.2, 4.5, 2.0, t["lingue_testo"], size=12.5)

    # --- 7 canali
    if x["canali"]:
        s = k.slide()
        k.titolo(s, "Maps e Ricerca", "Dove la campagna ha portato indicazioni stradali e visite")
        maps = next((c for c in x["canali"] if c["nome"] == "Google Maps"), {})
        ric = next((c for c in x["canali"] if c["nome"] == "Rete di Ricerca"), {})
        maps_p = next((c for c in x["canali_prec"] if c["nome"] == "Google Maps"), {})
        ric_p = next((c for c in x["canali_prec"] if c["nome"] == "Rete di Ricerca"), {})
        vlab = "Visite stimate" if stimate else "Visite"
        carte = [(num(maps.get("indicazioni")), "Indicazioni da Maps", f"{_var(maps.get('indicazioni'), maps_p.get('indicazioni'))} su {mese_prec}", True),
                 (num(maps.get("visite")), f"{vlab} da Maps",
                  f"{eur(maps['spesa'] / maps['visite'])} di spesa per visita" if maps.get("visite") else "", False),
                 (num(ric.get("indicazioni")), "Indicazioni da Ricerca", f"era {num(ric_p.get('indicazioni'))}", False),
                 (num(ric.get("visite")), f"{vlab} da Ricerca", f"era {num(ric_p.get('visite'))}", False)]
        for i, (v, e, nota, ev) in enumerate(carte):
            k.kpi(s, 0.6 + i * 3.1, 1.85, 2.9, 1.95, v, e, nota, ev)
        k.forma(s, 0.6, 4.1, 12.2, 1.9, "FFFFFF", bordo=T["bordo"])
        k.testo(s, 0.95, 4.3, 11.5, 1.6, [{"s": t["canali_titolo"], "size": 20, "bold": True, "font": "Georgia",
                                           "colore": T["primario"], "dopo": 8}, {"s": t["canali_testo"], "size": 13}])

    # --- 8 orari
    if sum(x["ore"]):
        s = k.slide()
        k.titolo(s, "Quando lavora la campagna", f"Conversioni per ora del giorno, {mese}")
        attive = [h for h in range(24) if x["ore"][h] > 0]
        ore = list(range(min(attive), max(attive) + 1))
        valori = [round(x["ore"][h]) for h in ore]
        top3 = sorted(range(len(ore)), key=lambda i: -valori[i])[:4]
        k.barre(s, 0.6, 1.7, 8.2, 4.9, [str(h) for h in ore], [("Conversioni", valori)], picco=set(top3), legenda=False)
        pranzo = sum(x["ore"][11:16]) / (sum(x["ore"]) or 1)
        k.forma(s, 9.2, 1.8, 3.6, 1.6, T["verde_bg"])
        k.testo(s, 9.2, 1.9, 3.6, 0.8, _pct(pranzo, 0), size=36, colore=T["verde"], bold=True, font="Georgia",
                align=PP_ALIGN.CENTER)
        k.testo(s, 9.35, 2.75, 3.3, 0.5, "delle conversioni a pranzo, dalle 11 alle 15", size=11, colore=T["verde"],
                align=PP_ALIGN.CENTER)
        giorni = sorted(range(7), key=lambda i: -x["giorni_conv"][i])
        righe = [["Giorno", "Conv.", "CPA"]] + [[GIORNI[i], num(x["giorni_conv"][i]),
                                                 eur(x["giorni_spesa"][i] / x["giorni_conv"][i]) if x["giorni_conv"][i] else "–"]
                                                for i in giorni]
        k.tabella(s, 9.2, 3.6, 3.6, righe, [1.6, 1, 1], altezza=0.36, size=11,
                  evidenzia={1: T["verde_bg"], len(righe) - 1: T["crema"]})

    # --- 9 economia
    comp = d["compenso"] or {}
    if d["mese_chiuso"] and x["coperti"] and x["scontrino"] and comp:
        s = k.slide()
        cop, sc = x["coperti"], x["scontrino"]
        if comp["tipo"] == "fisso":
            marg = comp.get("margine") or 0
            k.titolo(s, "Il ritorno economico", f"{_it(cop, 1)} coperti per visita, {eur(sc, 0)} a persona"
                     + (f", margine del {_it(marg * 100, 0)}%" if marg else "") + ", costi Google Ads e gestione inclusi")
            costi = cor["spesa"] + comp["importo"]

            def colonna(v):
                fat = v * cop * sc
                return [v, v * cop, fat, fat * marg, costi, fat * marg - costi]
            a, b = colonna(cor["visite"]), colonna(round(cor["visite"] * 1.3))
            righe = [["", "Dato provvisorio", "A dato consolidato (+30%)"],
                     ["Visite in negozio", num(a[0]), num(b[0])], ["Coperti stimati", num(a[1]), num(b[1])],
                     ["Fatturato attribuibile", eur(a[2], 0), eur(b[2], 0)],
                     [f"Margine generato ({_it(marg * 100, 0)}%)", eur(a[3], 0), eur(b[3], 0)],
                     ["Costi (ads e gestione)", "− " + eur(costi, 0), "− " + eur(costi, 0)],
                     ["Profitto netto", eur(a[5], 0), eur(b[5], 0)]]
            k.tabella(s, 0.6, 1.8, 8.0, righe, [3, 2, 2.4], evidenzia={6: T["crema"]})
            roi_a, roi_b = a[5] / costi, b[5] / costi
            k.forma(s, 9.0, 1.8, 3.8, 2.1, "FFFFFF", bordo=T["bordo"])
            k.testo(s, 9.0, 1.95, 3.8, 0.9, _pct(roi_a, 0), size=44, colore=T["primario"], bold=True, font="Georgia",
                    align=PP_ALIGN.CENTER)
            k.testo(s, 9.2, 2.9, 3.4, 0.8, f"ROI sul dato provvisorio, ROAS {_it(a[2] / costi, 0)} volte", size=12.5,
                    colore=T["tenue"], align=PP_ALIGN.CENTER)
            k.forma(s, 9.0, 4.1, 3.8, 2.1, T["verde_bg"])
            k.testo(s, 9.0, 4.25, 3.8, 0.9, _pct(roi_b, 0), size=44, colore=T["verde"], bold=True, font="Georgia",
                    align=PP_ALIGN.CENTER)
            k.testo(s, 9.2, 5.2, 3.4, 0.8, f"ROI atteso a dato consolidato, ROAS {_it(b[2] / costi, 0)} volte", size=12.5,
                    colore=T["verde"], align=PP_ALIGN.CENTER)
        else:
            perc = comp["percentuale"]
            k.titolo(s, "Fatturato attribuibile e compenso",
                     f"Come da accordo: {_it(perc * 100, 0)}% del fatturato attribuibile, {_it(cop, 1)} coperti per visita, "
                     f"{eur(sc, 0)} a persona")
            v = cor["visite"]
            fat = v * cop * sc
            righe = [["Voce", "Operazione", "Valore"]]
            if stimate:
                righe.append(["Indicazioni stradali", "dato Google", num(cor["indicazioni"])])
                righe.append(["Visite stimate", f"{num(cor['indicazioni'])} × {_it(x['quota_stima'] * 100, 0)}%", num(v)])
            else:
                righe.append(["Visite in negozio", "registrate a oggi", num(v)])
            righe += [["Coperti", f"{num(v)} × {_it(cop, 1)}", num(v * cop)],
                      ["Fatturato attribuibile", f"{num(v * cop)} × {eur(sc, 0)}", eur(fat, 0)],
                      [f"Compenso di {mese} ({_it(perc * 100, 0)}%)", f"{eur(fat, 0)} × {_it(perc * 100, 0)}%", eur(fat * perc, 0)]]
            k.tabella(s, 0.6, 1.8, 8.0, righe, [3, 2.4, 1.6], evidenzia={len(righe) - 2: T["verde_bg"], len(righe) - 1: T["crema"]})
            costo_tot = cor["spesa"] + fat * perc
            k.forma(s, 9.0, 4.1, 3.8, 2.1, T["verde_bg"])
            k.testo(s, 9.0, 4.25, 3.8, 0.9, f"{_it(fat / costo_tot, 1)} volte" if costo_tot else "–", size=36,
                    colore=T["verde"], bold=True, font="Georgia", align=PP_ALIGN.CENTER)
            k.testo(s, 9.2, 5.1, 3.4, 1.0, f"il fatturato attribuibile rispetto al costo totale per il locale "
                    f"({eur(cor['spesa'], 0)} di Google Ads più {eur(fat * perc, 0)} di compenso)", size=11,
                    colore=T["verde"], align=PP_ALIGN.CENTER)
            nota = "Le indicazioni stradali sono un dato definitivo: non è previsto conguaglio." if stimate else (
                "Le visite registrate dopo oggi saranno conteggiate a conguaglio nella fattura successiva.")
            k.testo(s, 0.6, 1.9 + 0.36 * len(righe) + 0.2, 8, 0.5, nota, size=11, colore=T["tenue"], italic=True)

    # --- 10 prossimo mese
    s = k.slide(scura=True)
    k.forma(s, 10.2, 4.6, 4.6, 4.6, T["primario"], MSO_SHAPE.OVAL)
    k.testo(s, 0.9, 0.7, 11, 0.9, f"{d['mese_succ'].capitalize()}: cosa facciamo", size=34, colore="FFFFFF", bold=True,
            font="Georgia")
    for i, (tit, testo) in enumerate(t["prossimi"][:4]):
        y = 1.85 + i * 1.2
        k.testo(s, 0.9, y, 1, 0.6, f"{i + 1:02d}", size=30, colore=T["oro"], bold=True, font="Georgia")
        k.testo(s, 2.0, y + 0.02, 9, 0.45, tit, size=19, colore="FFFFFF", bold=True)
        k.testo(s, 2.0, y + 0.48, 9.5, 0.5, testo, size=14, colore=T["cover_txt"])
    return k.salva(percorso)
